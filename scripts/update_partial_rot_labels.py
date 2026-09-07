from __future__ import annotations

import argparse
import hashlib
import json
import os
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def zip_image_ids(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return sorted(
            Path(name).stem
            for name in archive.namelist()
            if name and not name.endswith("/")
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Merge partial manual rot labels into the review workbook."
    )
    parser.add_argument("--source-workbook", type=Path, required=True)
    parser.add_argument("--part1-zip", type=Path, required=True)
    parser.add_argument("--part1-positive-dir", type=Path)
    parser.add_argument("--part2-zip", type=Path, required=True)
    parser.add_argument("--part2-labels", type=Path, required=True)
    parser.add_argument("--part3-zip", type=Path, required=True)
    parser.add_argument("--part3-positive-dir", type=Path, required=True)
    parser.add_argument("--output-workbook", type=Path, required=True)
    parser.add_argument("--audit-json", type=Path, required=True)
    args = parser.parse_args()

    source_hash_before = sha256(args.source_workbook)
    part1_ids = zip_image_ids(args.part1_zip)
    part2_ids = zip_image_ids(args.part2_zip)
    part3_ids = zip_image_ids(args.part3_zip)
    part2_labels = [
        line.strip().lower()
        for line in args.part2_labels.read_text(encoding="utf-8-sig").splitlines()
        if line.strip()
    ]
    invalid = [label for label in part2_labels if label not in {"y", "n"}]
    if invalid:
        raise ValueError(f"Invalid labels in part 2: {sorted(set(invalid))}")
    if len(part2_ids) != len(part2_labels):
        raise ValueError(
            f"Part 2 count mismatch: {len(part2_ids)} images vs "
            f"{len(part2_labels)} labels"
        )

    part1_positive_ids: set[str] = set()
    if args.part1_positive_dir is not None:
        part1_positive_ids = {
            path.stem for path in args.part1_positive_dir.iterdir() if path.is_file()
        }
        unexpected_part1_ids = sorted(part1_positive_ids - set(part1_ids))
        if unexpected_part1_ids:
            raise ValueError(
                "Part 1 positive labels not present in part 1 archive: "
                + ", ".join(unexpected_part1_ids[:10])
            )

    part3_positive_ids = {
        path.stem for path in args.part3_positive_dir.glob("*.txt") if path.is_file()
    }
    unexpected_positive_ids = sorted(part3_positive_ids - set(part3_ids))
    if unexpected_positive_ids:
        raise ValueError(
            "Part 3 positive labels not present in part 3 archive: "
            + ", ".join(unexpected_positive_ids[:10])
        )

    groups = [set(part1_ids), set(part2_ids), set(part3_ids)]
    if groups[0] & groups[1] or groups[0] & groups[2] or groups[1] & groups[2]:
        raise ValueError("The three archive rosters overlap")

    if args.part1_positive_dir is None:
        merged_labels: dict[str, str | None] = {
            instance_id: None for instance_id in part1_ids
        }
    else:
        merged_labels = {
            instance_id: "y" if instance_id in part1_positive_ids else "n"
            for instance_id in part1_ids
        }
    merged_labels.update(dict(zip(part2_ids, part2_labels, strict=True)))
    merged_labels.update(
        {
            instance_id: "y" if instance_id in part3_positive_ids else "n"
            for instance_id in part3_ids
        }
    )

    workbook = load_workbook(args.source_workbook)
    if "rot_review" not in workbook.sheetnames:
        raise ValueError("Worksheet 'rot_review' not found")
    sheet = workbook["rot_review"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    if "instance_id" not in headers or "rot" not in headers:
        raise ValueError("Required columns 'instance_id' and 'rot' not found")

    instance_col = headers["instance_id"]
    rot_col = headers["rot"]
    workbook_ids: list[str] = []
    before_counts: Counter[str] = Counter()
    after_counts: Counter[str] = Counter()
    for row in range(2, sheet.max_row + 1):
        instance_id = str(sheet.cell(row=row, column=instance_col).value or "").strip()
        if not instance_id:
            raise ValueError(f"Missing instance_id at row {row}")
        workbook_ids.append(instance_id)
        old_value = sheet.cell(row=row, column=rot_col).value
        before_counts[str(old_value).strip().lower() if old_value is not None else "blank"] += 1
        if instance_id not in merged_labels:
            raise ValueError(f"Workbook instance is absent from the three rosters: {instance_id}")
        new_value = merged_labels[instance_id]
        sheet.cell(row=row, column=rot_col).value = new_value
        after_counts[new_value if new_value is not None else "blank"] += 1

    workbook_id_set = set(workbook_ids)
    if len(workbook_ids) != len(workbook_id_set):
        raise ValueError("Duplicate instance_id values in workbook")
    roster_ids = set(merged_labels)
    if workbook_id_set != roster_ids:
        missing = sorted(roster_ids - workbook_id_set)
        extra = sorted(workbook_id_set - roster_ids)
        raise ValueError(f"Roster/workbook mismatch; missing={missing[:5]}, extra={extra[:5]}")

    args.output_workbook.parent.mkdir(parents=True, exist_ok=True)
    same_file = args.output_workbook.resolve() == args.source_workbook.resolve()
    save_target = (
        args.output_workbook.with_name(args.output_workbook.stem + ".tmp.xlsx")
        if same_file
        else args.output_workbook
    )
    workbook.save(save_target)
    workbook.close()
    if same_file:
        os.replace(save_target, args.output_workbook)

    # Re-open the saved file and reconcile every edited value.
    verified = load_workbook(args.output_workbook, read_only=True, data_only=False)
    verify_sheet = verified["rot_review"]
    verify_headers = {
        cell.value: cell.column for cell in next(verify_sheet.iter_rows(min_row=1, max_row=1))
    }
    verify_counts: Counter[str] = Counter()
    mismatch_count = 0
    formula_error_count = 0
    formula_errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
    for row in verify_sheet.iter_rows(min_row=2):
        instance_id = str(row[verify_headers["instance_id"] - 1].value or "").strip()
        value = row[verify_headers["rot"] - 1].value
        normalized = str(value).strip().lower() if value is not None else "blank"
        verify_counts[normalized] += 1
        expected = merged_labels[instance_id]
        if value != expected:
            mismatch_count += 1
        for cell in row:
            if isinstance(cell.value, str) and cell.value in formula_errors:
                formula_error_count += 1
    verified.close()

    if mismatch_count:
        raise ValueError(f"Saved workbook has {mismatch_count} label mismatches")
    if not same_file and source_hash_before != sha256(args.source_workbook):
        raise ValueError("Source workbook changed during processing")

    audit = {
        "source_workbook": str(args.source_workbook),
        "source_sha256": source_hash_before,
        "output_workbook": str(args.output_workbook),
        "output_sha256": sha256(args.output_workbook),
        "worksheet": "rot_review",
        "edited_column": "rot",
        "row_count": len(workbook_ids),
        "part1": {
            "count": len(part1_ids),
            "range": [part1_ids[0], part1_ids[-1]],
            "status": "complete" if args.part1_positive_dir is not None else "blank_pending",
            "y": len(part1_positive_ids) if args.part1_positive_dir is not None else None,
            "n": (
                len(part1_ids) - len(part1_positive_ids)
                if args.part1_positive_dir is not None
                else None
            ),
        },
        "part2": {
            "count": len(part2_ids),
            "range": [part2_ids[0], part2_ids[-1]],
            "y": part2_labels.count("y"),
            "n": part2_labels.count("n"),
        },
        "part3": {
            "count": len(part3_ids),
            "range": [part3_ids[0], part3_ids[-1]],
            "y": len(part3_positive_ids),
            "n": len(part3_ids) - len(part3_positive_ids),
        },
        "before_counts": dict(before_counts),
        "after_counts": dict(after_counts),
        "verified_counts": dict(verify_counts),
        "mismatch_count": mismatch_count,
        "formula_error_count": formula_error_count,
    }
    args.audit_json.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
