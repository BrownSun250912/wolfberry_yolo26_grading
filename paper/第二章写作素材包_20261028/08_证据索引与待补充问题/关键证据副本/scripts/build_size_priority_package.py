from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
from pathlib import Path

import cv2
import matplotlib as mpl

mpl.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from build_grading_scheme_candidates import COLOR_FEATURE_WEIGHTS, add_component_scores


RANDOM_SEED = 20260902
SCHEME_NAME = "size_priority"
SCHEME_WEIGHTS = {
    "size_score": 0.55,
    "shape_score": 0.30,
    "color_uniformity_score": 0.15,
}
LOWER_QUANTILE = 0.30
UPPER_QUANTILE = 0.70
BOUNDARY_SAMPLE_PER_SIDE = 15


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Freeze size-priority candidate grading and build a boundary review pack.")
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument(
        "--features-csv",
        type=Path,
        default=Path("outputs/strict_rot_exclusion_v1/qualified_features.csv"),
    )
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/size_priority_v1"))
    parser.add_argument("--seed", type=int, default=RANDOM_SEED)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_reference_values(data: pd.DataFrame) -> pd.DataFrame:
    train = data.loc[data["split"].eq("train")]
    rows: list[dict[str, object]] = []

    for feature in ("area", "minor_axis", "solidity", "aspect_ratio", "circularity"):
        values = np.sort(train[feature].to_numpy(dtype=float))
        for rank, value in enumerate(values, start=1):
            rows.append(
                {
                    "reference_scope": "global_train",
                    "source_group": "ALL",
                    "feature": feature,
                    "rank": rank,
                    "n": len(values),
                    "value": float(value),
                }
            )

    for group in sorted(data["source_group"].astype(str).unique()):
        group_train = train.loc[train["source_group"].astype(str).eq(group)]
        for feature in COLOR_FEATURE_WEIGHTS:
            if len(group_train) >= 30:
                values = np.sort(group_train[feature].to_numpy(dtype=float))
                scope = "source_group_train"
                reference_group = group
            else:
                values = np.sort(train[feature].to_numpy(dtype=float))
                scope = "global_train_fallback"
                reference_group = "ALL"
            for rank, value in enumerate(values, start=1):
                rows.append(
                    {
                        "reference_scope": scope,
                        "source_group": reference_group,
                        "feature": feature,
                        "rank": rank,
                        "n": len(values),
                        "value": float(value),
                    }
                )
    return pd.DataFrame(rows)


def select_side(
    pool: pd.DataFrame,
    count: int,
    used_instances: set[str],
    used_images: set[str],
) -> pd.DataFrame:
    available = pool.loc[~pool["instance_id"].astype(str).isin(used_instances)].sort_values(
        ["absolute_distance_to_boundary", "instance_id"], ascending=[True, True]
    )
    selected: list[pd.Series] = []
    selected_ids: set[str] = set()
    for _, row in available.iterrows():
        image_id = str(row["image_id"])
        instance_id = str(row["instance_id"])
        if image_id in used_images:
            continue
        selected.append(row)
        selected_ids.add(instance_id)
        used_instances.add(instance_id)
        used_images.add(image_id)
        if len(selected) == count:
            break
    if len(selected) < count:
        for _, row in available.iterrows():
            instance_id = str(row["instance_id"])
            if instance_id in selected_ids:
                continue
            selected.append(row)
            selected_ids.add(instance_id)
            used_instances.add(instance_id)
            if len(selected) == count:
                break
    if len(selected) != count:
        raise RuntimeError(f"Could only select {len(selected)} of {count} requested boundary samples.")
    return pd.DataFrame(selected)


def select_boundary_review(scored: pd.DataFrame, low: float, high: float) -> pd.DataFrame:
    used_instances: set[str] = set()
    used_images: set[str] = set()
    groups: list[pd.DataFrame] = []

    definitions = [
        ("boundary_grade1_2", "grade1_side", high, scored["candidate_grade"].eq(1)),
        ("boundary_grade1_2", "grade2_side", high, scored["candidate_grade"].eq(2)),
        ("boundary_grade2_3", "grade2_side", low, scored["candidate_grade"].eq(2)),
        ("boundary_grade2_3", "grade3_side", low, scored["candidate_grade"].eq(3)),
    ]
    for group_name, side_name, boundary, condition in definitions:
        pool = scored.loc[condition].copy()
        pool["candidate_group"] = group_name
        pool["boundary_side"] = side_name
        pool["boundary_score"] = boundary
        pool["signed_distance_to_boundary"] = pool["size_priority_score"] - boundary
        pool["absolute_distance_to_boundary"] = pool["signed_distance_to_boundary"].abs()
        groups.append(select_side(pool, BOUNDARY_SAMPLE_PER_SIDE, used_instances, used_images))

    candidates = pd.concat(groups, ignore_index=True)
    side_order = {"grade1_side": 1, "grade2_side": 2, "grade3_side": 3}
    candidates["_side_order"] = candidates["boundary_side"].map(side_order)
    candidates = candidates.sort_values(
        ["candidate_group", "_side_order", "absolute_distance_to_boundary", "instance_id"]
    ).drop(columns="_side_order")
    candidates = candidates.reset_index(drop=True)
    candidates.insert(0, "review_order", np.arange(1, len(candidates) + 1))
    candidates.insert(0, "review_id", [f"SP{index:03d}" for index in range(1, len(candidates) + 1)])
    candidates["review_image_file"] = candidates["review_id"] + "__" + candidates["instance_id"].astype(str) + ".png"
    candidates["review_expected_grade"] = ""
    candidates["review_accept_candidate"] = ""
    candidates["review_comment"] = ""
    return candidates


def place_image(canvas: np.ndarray, source: np.ndarray, left: int, top: int, width: int, height: int) -> None:
    height_source, width_source = source.shape[:2]
    scale = min(width / max(width_source, 1), height / max(height_source, 1))
    resized = cv2.resize(
        source,
        (max(1, int(round(width_source * scale))), max(1, int(round(height_source * scale)))),
        interpolation=cv2.INTER_AREA,
    )
    rh, rw = resized.shape[:2]
    y = top + (height - rh) // 2
    x = left + (width - rw) // 2
    canvas[y : y + rh, x : x + rw] = resized


def make_review_image(row: pd.Series, project_root: Path, output_path: Path) -> None:
    image = cv2.imread(str(Path(str(row["image_path"]))), cv2.IMREAD_COLOR)
    artifact_root = project_root / "outputs" / "features_frozen_v1" / str(row["split"])
    crop = cv2.imread(str(artifact_root / str(row["crop_path"])), cv2.IMREAD_UNCHANGED)
    mask = cv2.imread(str(artifact_root / str(row["mask_path"])), cv2.IMREAD_GRAYSCALE)
    if image is None or crop is None or mask is None:
        raise FileNotFoundError(f"Missing review artifacts for {row['instance_id']}")

    x1, y1, x2, y2 = [int(row[name]) for name in ("roi_x1", "roi_y1", "roi_x2", "roi_y2")]
    original_roi = image[y1:y2, x1:x2].copy()
    contours, _ = cv2.findContours((mask > 0).astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    outlined = original_roi.copy()
    cv2.drawContours(outlined, contours, -1, (0, 190, 0), 2)

    if crop.ndim == 3 and crop.shape[2] == 4:
        alpha = crop[:, :, 3:4].astype(np.float32) / 255.0
        isolated = (crop[:, :, :3].astype(np.float32) * alpha + 32.0 * (1.0 - alpha)).astype(np.uint8)
    else:
        isolated = crop[:, :, :3]

    canvas = np.full((430, 1120, 3), 247, dtype=np.uint8)
    title = f"{row['review_id']}  {row['instance_id']}  {row['candidate_group']} / {row['boundary_side']}"
    cv2.putText(canvas, title, (18, 28), cv2.FONT_HERSHEY_SIMPLEX, 0.58, (18, 18, 18), 2, cv2.LINE_AA)
    cv2.putText(canvas, "Refined mask contour", (20, 55), cv2.FONT_HERSHEY_SIMPLEX, 0.48, (45, 45, 45), 1)
    cv2.putText(canvas, "Background-excluded crop", (380, 55), cv2.FONT_HERSHEY_SIMPLEX, 0.48, (45, 45, 45), 1)
    place_image(canvas, outlined, 20, 65, 320, 320)
    place_image(canvas, isolated, 380, 65, 320, 320)

    delta = float(row["signed_distance_to_boundary"])
    lines = [
        f"Candidate grade: {int(row['candidate_grade'])}",
        f"Size-priority score: {row['size_priority_score']:.2f}",
        f"Boundary: {row['boundary_score']:.2f}  delta={delta:+.2f}",
        f"Size / shape / color: {row['size_score']:.1f} / {row['shape_score']:.1f} / {row['color_uniformity_score']:.1f}",
        f"Area={row['area']:.0f} px2  minor={row['minor_axis']:.1f} px",
        f"Aspect={row['aspect_ratio']:.2f}  solidity={row['solidity']:.3f}",
        f"Circularity={row['circularity']:.3f}",
        f"Lstd={row['L_std']:.2f}  astd={row['a_std']:.2f}  bstd={row['b_std']:.2f}",
        f"max_damage_ratio={row['max_damage_ratio']:.3f}",
        f"Batch={row['source_group']}  split={row['split']}",
        "Candidate only; confirm before training use.",
    ]
    y = 76
    for index, line in enumerate(lines):
        color = (15, 15, 15) if index < 4 else (45, 45, 45)
        thickness = 2 if index == 0 else 1
        cv2.putText(canvas, line, (725, y), cv2.FONT_HERSHEY_SIMPLEX, 0.45, color, thickness, cv2.LINE_AA)
        y += 30

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not cv2.imwrite(str(output_path), canvas):
        raise IOError(output_path)


def make_contact_sheets(candidates: pd.DataFrame, review_dir: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for group, group_frame in candidates.groupby("candidate_group", sort=False):
        cols = 5
        rows = math.ceil(len(group_frame) / cols)
        tile_w, tile_h = 400, 154
        title_h = 58
        sheet = np.full((title_h + rows * tile_h, cols * tile_w, 3), 250, dtype=np.uint8)
        title = f"size_priority: {group}  n={len(group_frame)} (15 per side)"
        cv2.putText(sheet, title, (18, 38), cv2.FONT_HERSHEY_SIMPLEX, 0.82, (20, 20, 20), 2, cv2.LINE_AA)
        for index, (_, row) in enumerate(group_frame.iterrows()):
            panel = cv2.imread(str(review_dir / str(row["review_image_file"])), cv2.IMREAD_COLOR)
            if panel is None:
                raise FileNotFoundError(row["review_image_file"])
            panel = cv2.resize(panel, (tile_w, tile_h), interpolation=cv2.INTER_AREA)
            top = title_h + (index // cols) * tile_h
            left = (index % cols) * tile_w
            sheet[top : top + tile_h, left : left + tile_w] = panel
        output_path = output_dir / f"{group}_contact_sheet.png"
        if not cv2.imwrite(str(output_path), sheet):
            raise IOError(output_path)


def make_score_distribution(scored: pd.DataFrame, low: float, high: float, output_path: Path) -> None:
    mpl.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )
    fig, ax = plt.subplots(figsize=(9.2, 5.6), layout="constrained")
    bins = np.linspace(0, 100, 41)
    ax.hist(scored["size_priority_score"], bins=bins, color="#6BAED6", edgecolor="white", linewidth=0.5)
    ax.axvline(low, color="#D55E00", linestyle="--", linewidth=2, label=f"Grade 3/2 candidate boundary = {low:.2f}")
    ax.axvline(high, color="#0072B2", linestyle="-.", linewidth=2, label=f"Grade 2/1 candidate boundary = {high:.2f}")
    ax.text(low - 1.0, ax.get_ylim()[1] * 0.91, "G3 | G2", ha="right", va="top", color="#8C3B00")
    ax.text(high + 1.0, ax.get_ylim()[1] * 0.91, "G2 | G1", ha="left", va="top", color="#004E7A")
    ax.set(
        title="Size-priority candidate score distribution",
        xlabel="Size-priority score (0-100; empirical-percentile composite)",
        ylabel="Instance count (40 fixed bins)",
        xlim=(0, 100),
    )
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, loc="upper left")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def make_distribution_table(scored: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for split in ("all", "train", "val"):
        subset = scored if split == "all" else scored.loc[scored["split"].eq(split)]
        for grade in (1, 2, 3):
            count = int(subset["candidate_grade"].eq(grade).sum())
            rows.append(
                {
                    "split": split,
                    "candidate_grade": grade,
                    "count": count,
                    "share": count / len(subset),
                    "label_status": "candidate_only_not_training_truth",
                }
            )
    return pd.DataFrame(rows)


def bootstrap_thresholds(scored: pd.DataFrame, seed: int, iterations: int = 500) -> pd.DataFrame:
    values = scored.loc[scored["split"].eq("train"), "size_priority_score"].to_numpy(dtype=float)
    rng = np.random.default_rng(seed)
    output: list[dict[str, object]] = []
    for boundary, quantile in (("grade3_to_2", LOWER_QUANTILE), ("grade2_to_1", UPPER_QUANTILE)):
        estimates = np.empty(iterations, dtype=float)
        for iteration in range(iterations):
            estimates[iteration] = np.quantile(rng.choice(values, size=len(values), replace=True), quantile)
        output.append(
            {
                "scheme": SCHEME_NAME,
                "boundary": boundary,
                "quantile": quantile,
                "bootstrap_iterations": iterations,
                "mean": float(np.mean(estimates)),
                "std": float(np.std(estimates, ddof=1)),
                "P2_5": float(np.quantile(estimates, 0.025)),
                "median": float(np.quantile(estimates, 0.50)),
                "P97_5": float(np.quantile(estimates, 0.975)),
                "random_seed": seed,
            }
        )
    return pd.DataFrame(output)


def write_review_workbook(review: pd.DataFrame, output_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BoundaryReview"
    instructions = workbook.create_sheet("Instructions")

    instruction_rows = [
        ["size_priority_v1 边界审核说明"],
        ["用途", "审核候选分界附近样本；不是重新判断腐烂，也不是最终训练标签。"],
        ["review_expected_grade", "填写你认为正确的 1、2 或 3。"],
        ["review_accept_candidate", "候选等级可接受填 YES，否则填 NO。"],
        ["review_comment", "可记录太小、太细、形态不规则、颜色不均或边界应上移/下移。"],
        ["图片对应", "review_image_file 与 review_images 文件夹一一对应，工作表内可点击。"],
        ["分界抽样", "1/2 与 2/3 两条分界各 30 颗；每条分界两侧各 15 颗。"],
        ["重要限制", "尺寸单位仍是像素；没有毫米标定时，只能在拍摄距离与缩放一致的数据中比较。"],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    instructions["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    instructions.merge_cells("A1:B1")
    instructions.column_dimensions["A"].width = 26
    instructions.column_dimensions["B"].width = 90
    for row in instructions.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="1F4E78")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    instructions.sheet_view.showGridLines = False

    for column_index, name in enumerate(review.columns, start=1):
        sheet.cell(1, column_index, name)
    for row_index, values in enumerate(review.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(values, start=1):
            if pd.isna(value):
                value = None
            sheet.cell(row_index, column_index, value)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 36
    sheet.freeze_panes = "G2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False

    grade_column = review.columns.get_loc("candidate_grade") + 1
    image_column = review.columns.get_loc("review_image_file") + 1
    expected_column = review.columns.get_loc("review_expected_grade") + 1
    accept_column = review.columns.get_loc("review_accept_candidate") + 1
    comment_column = review.columns.get_loc("review_comment") + 1
    for row_index in range(2, len(review) + 2):
        image_cell = sheet.cell(row_index, image_column)
        image_cell.hyperlink = f"review_images/{image_cell.value}"
        image_cell.style = "Hyperlink"
        for column_index in (expected_column, accept_column, comment_column):
            sheet.cell(row_index, column_index).fill = editable_fill

    expected_letter = openpyxl.utils.get_column_letter(expected_column)
    accept_letter = openpyxl.utils.get_column_letter(accept_column)
    grade_letter = openpyxl.utils.get_column_letter(grade_column)
    grade_validation = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    accept_validation = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    sheet.add_data_validation(grade_validation)
    sheet.add_data_validation(accept_validation)
    grade_validation.add(f"{expected_letter}2:{expected_letter}{len(review)+1}")
    accept_validation.add(f"{accept_letter}2:{accept_letter}{len(review)+1}")

    grade_fills = {1: "C6E0B4", 2: "FFE699", 3: "F4B084"}
    for grade, color in grade_fills.items():
        sheet.conditional_formatting.add(
            f"{grade_letter}2:{grade_letter}{len(review)+1}",
            FormulaRule(formula=[f"{grade_letter}2={grade}"], fill=PatternFill("solid", fgColor=color)),
        )

    widths = {
        "review_id": 11,
        "review_order": 11,
        "candidate_group": 22,
        "boundary_side": 16,
        "image_id": 14,
        "instance_id": 22,
        "candidate_grade": 14,
        "review_image_file": 32,
        "review_expected_grade": 20,
        "review_accept_candidate": 22,
        "review_comment": 42,
    }
    for index, column in enumerate(review.columns, start=1):
        width = widths.get(column, 15)
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        if pd.api.types.is_float_dtype(review[column]):
            for row_index in range(2, len(review) + 2):
                sheet.cell(row_index, index).number_format = "0.000"

    table = Table(displayName="SizePriorityBoundaryReview", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)
    workbook.save(output_path)


def write_readme(output_dir: Path, total: int, train_count: int, val_count: int, low: float, high: float) -> None:
    text = f"""# size_priority_v1 枸杞三级候选方案

本目录基于严格人工排除腐烂后的 {total:,} 颗实例生成（train={train_count:,}，val={val_count:,}）。这里的 1、2、3 级仍是**候选标签**，尚未经过边界人工审核，不能直接当作最终商品标准或最终训练真值。

## 评分方法

- 尺寸分数：`65% × area 的 train 经验百分位 + 35% × minor_axis 的 train 经验百分位`。
- 形态分数：`50% × solidity 百分位 + 30% × aspect_ratio 典型度 + 20% × circularity 百分位`。
- 颜色均匀性分数：对 `L_std、a_std、b_std、max_damage_ratio` 做反向经验百分位后，按 `20%、30%、20%、30%` 加权。颜色参考分布按 `source_group` 的 train 数据拟合，用于减弱分几天拍摄造成的灯光批次差异。
- size_priority 总分：`55% × 尺寸 + 30% × 形态 + 15% × 颜色均匀性`。

尺寸与形态参考分布只使用全局 train 数据，不按拍摄批次归一化，以免把真实大小差异消掉。全部经验分布原值保存在 `scoring_reference_values.csv`，便于复现。

## 当前候选分界

- 分数 `>= {high:.6f}`：候选 1 级。
- 分数 `>= {low:.6f}` 且 `< {high:.6f}`：候选 2 级。
- 分数 `< {low:.6f}`：候选 3 级。

两条分界来自 **train 集 size_priority 分数的 P70 和 P30**，不是人工确认后的最终阈值。

## 人工审核

优先打开 `boundary_review.xlsx`。60 颗样本中：

- `boundary_grade1_2`：30 颗，分界两侧各 15 颗。
- `boundary_grade2_3`：30 颗，分界两侧各 15 颗。

填写黄色列：

- `review_expected_grade`：你认为应为 1、2、3 中哪一级。
- `review_accept_candidate`：候选等级可接受填 `YES`，否则填 `NO`。
- `review_comment`：写清太小、太细、形态不规则、颜色不均或分界应调整的原因。

工作表中的 `review_image_file` 可点击打开对应图片。`contact_sheets/` 用于快速总览，但细看时以单张审核图为准。

## 文件说明

- `size_priority_candidate_labels.csv`：全部 {total:,} 颗的分项分数、总分和候选等级。
- `boundary_review.csv` / `boundary_review.xlsx`：60 颗边界审核索引。
- `size_priority_config.json`：冻结公式、权重、分界、输入哈希和版本信息。
- `scoring_reference_values.csv`：拟合于 train 的完整经验参考分布。
- `score_reference_quantiles.csv`：参考分布摘要。
- `grade_distribution.csv`：train、val 和全量候选等级数量。
- `threshold_bootstrap_stability.csv`：500 次 bootstrap 的分界稳定性。
- `plots/size_priority_score_distribution.png`：总分分布与候选分界。

## 限制

当前尺寸单位为像素，没有毫米标定。只有在相机距离、缩放和拍摄布局基本一致时，像素尺寸才适合直接比较；若未来加入不同距离或不同分辨率的新数据，需要重新标定或引入尺度参照物。
"""
    (output_dir / "README.md").write_text(text, encoding="utf-8")


def main() -> int:
    args = parse_args()
    project_root = args.project_root.resolve()
    features_path = args.features_csv if args.features_csv.is_absolute() else project_root / args.features_csv
    output_dir = args.output_dir if args.output_dir.is_absolute() else project_root / args.output_dir
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"Output directory is not empty: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    data = pd.read_csv(features_path, encoding="utf-8-sig")
    required = {
        "image_id",
        "instance_id",
        "split",
        "source_group",
        "image_path",
        "crop_path",
        "mask_path",
        "roi_x1",
        "roi_y1",
        "roi_x2",
        "roi_y2",
        "area",
        "minor_axis",
        "solidity",
        "aspect_ratio",
        "circularity",
        *COLOR_FEATURE_WEIGHTS,
    }
    missing = sorted(required - set(data.columns))
    if missing:
        raise KeyError(f"Missing required columns: {missing}")
    if data["instance_id"].duplicated().any():
        raise ValueError("instance_id must be unique.")
    feature_columns = sorted(required - {"image_id", "instance_id", "split", "source_group", "image_path", "crop_path", "mask_path"})
    if data[feature_columns].isna().any().any():
        raise ValueError("Required values contain missing data.")

    scored, reference_quantiles = add_component_scores(data)
    scored["size_priority_score"] = sum(scored[column] * weight for column, weight in SCHEME_WEIGHTS.items())
    train = scored["split"].eq("train")
    low = float(scored.loc[train, "size_priority_score"].quantile(LOWER_QUANTILE))
    high = float(scored.loc[train, "size_priority_score"].quantile(UPPER_QUANTILE))
    scored["candidate_grade"] = np.where(
        scored["size_priority_score"].ge(high),
        1,
        np.where(scored["size_priority_score"].lt(low), 3, 2),
    ).astype(int)
    scored["candidate_label_status"] = "candidate_only_not_training_truth"
    scored["grading_scheme_version"] = "size_priority_v1"

    review = select_boundary_review(scored, low, high)
    review_dir = output_dir / "review_images"
    for _, row in review.iterrows():
        make_review_image(row, project_root, review_dir / str(row["review_image_file"]))
    make_contact_sheets(review, review_dir, output_dir / "contact_sheets")
    make_score_distribution(scored, low, high, output_dir / "plots" / "size_priority_score_distribution.png")

    references = build_reference_values(data)
    distributions = make_distribution_table(scored)
    bootstrap = bootstrap_thresholds(scored, args.seed)

    preferred_columns = [
        "image_id",
        "instance_id",
        "split",
        "source_group",
        "grading_scheme_version",
        "candidate_label_status",
        "candidate_grade",
        "size_priority_score",
        "size_score",
        "shape_score",
        "color_uniformity_score",
        "pct_area",
        "pct_minor_axis",
        "pct_solidity",
        "pct_aspect_ratio",
        "aspect_typicality",
        "pct_circularity",
        *[f"uniformity_{feature}" for feature in COLOR_FEATURE_WEIGHTS],
    ]
    remaining_columns = [column for column in scored.columns if column not in preferred_columns]
    scored[preferred_columns + remaining_columns].to_csv(
        output_dir / "size_priority_candidate_labels.csv", index=False, encoding="utf-8-sig"
    )
    reference_quantiles.to_csv(output_dir / "score_reference_quantiles.csv", index=False, encoding="utf-8-sig")
    references.to_csv(output_dir / "scoring_reference_values.csv", index=False, encoding="utf-8-sig")
    distributions.to_csv(output_dir / "grade_distribution.csv", index=False, encoding="utf-8-sig")
    bootstrap.to_csv(output_dir / "threshold_bootstrap_stability.csv", index=False, encoding="utf-8-sig")

    review_columns = [
        "review_id",
        "review_order",
        "candidate_group",
        "boundary_side",
        "image_id",
        "instance_id",
        "candidate_grade",
        "size_priority_score",
        "boundary_score",
        "signed_distance_to_boundary",
        "absolute_distance_to_boundary",
        "size_score",
        "shape_score",
        "color_uniformity_score",
        "area",
        "minor_axis",
        "aspect_ratio",
        "solidity",
        "circularity",
        "L_std",
        "a_std",
        "b_std",
        "max_damage_ratio",
        "source_group",
        "split",
        "review_image_file",
        "review_expected_grade",
        "review_accept_candidate",
        "review_comment",
    ]
    review_export = review[review_columns].copy()
    review_export.to_csv(output_dir / "boundary_review.csv", index=False, encoding="utf-8-sig")
    write_review_workbook(review_export, output_dir / "boundary_review.xlsx")

    config = {
        "schema_version": 1,
        "scheme_name": SCHEME_NAME,
        "scheme_version": "size_priority_v1",
        "status": "candidate_only_not_training_truth",
        "input": {
            "features_csv": str(features_path.relative_to(project_root)).replace("\\", "/"),
            "sha256": sha256_file(features_path),
            "instances": int(len(scored)),
            "train_instances": int(train.sum()),
            "val_instances": int((~train).sum()),
            "manual_rot_exclusion": "strict second-reviewer labels already applied upstream",
        },
        "component_formulas": {
            "size_score": "100 * (0.65 * ECDF_train(area) + 0.35 * ECDF_train(minor_axis))",
            "aspect_typicality": "clip(1 - 2 * abs(ECDF_train(aspect_ratio) - 0.5), 0, 1)",
            "shape_score": "100 * (0.50 * ECDF_train(solidity) + 0.30 * aspect_typicality + 0.20 * ECDF_train(circularity))",
            "color_uniformity_score": "100 * sum(weight * (1 - ECDF_source_group_train(feature)))",
            "color_feature_weights": COLOR_FEATURE_WEIGHTS,
            "color_batch_policy": "source_group train reference; global train fallback if group train n < 30",
            "ecdf_tie_policy": "midrank = (left_search_rank + right_search_rank) / (2*n)",
        },
        "overall_formula": "0.55 * size_score + 0.30 * shape_score + 0.15 * color_uniformity_score",
        "overall_weights": SCHEME_WEIGHTS,
        "threshold_fit_partition": "train only",
        "threshold_quantiles": {"grade3_to_2": LOWER_QUANTILE, "grade2_to_1": UPPER_QUANTILE},
        "thresholds": {"grade3_to_2": low, "grade2_to_1": high},
        "grade_rule": {
            "grade_1": f"score >= {high}",
            "grade_2": f"{low} <= score < {high}",
            "grade_3": f"score < {low}",
        },
        "review_sampling": {
            "total": int(len(review)),
            "per_boundary": 2 * BOUNDARY_SAMPLE_PER_SIDE,
            "per_boundary_side": BOUNDARY_SAMPLE_PER_SIDE,
            "ordering": "absolute distance to boundary, then instance_id",
            "unique_image_preference": True,
            "random_seed": args.seed,
        },
        "physical_scale": {
            "calibrated": False,
            "size_units": "pixel and pixel^2",
            "limitation": "Valid for relative comparison only when camera distance and image scaling are consistent.",
        },
        "software": {
            "python": platform.python_version(),
            "pandas": pd.__version__,
            "numpy": np.__version__,
            "opencv": cv2.__version__,
            "matplotlib": mpl.__version__,
            "openpyxl": openpyxl.__version__,
        },
    }
    (output_dir / "size_priority_config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_readme(output_dir, len(scored), int(train.sum()), int((~train).sum()), low, high)

    summary = {
        "output_dir": str(output_dir),
        "instances": int(len(scored)),
        "train": int(train.sum()),
        "val": int((~train).sum()),
        "threshold_grade3_to_2": low,
        "threshold_grade2_to_1": high,
        "grade_counts_all": scored["candidate_grade"].value_counts().sort_index().to_dict(),
        "review_instances": int(len(review)),
        "review_unique_instances": int(review["instance_id"].nunique()),
        "review_unique_images": int(review["image_id"].nunique()),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
